#!/usr/bin/env python3
"""
Rebuild subject JSON data from the 2026 Excel curriculum guide.

Reads:
- 2026학년도 입학생을 위한 2022 개정 교육과정 선택 과목 안내서(EXCEL).xlsm
- app/src/data/json/school.json

Writes:
- data/subjects.json
- app/src/data/json/subjects.json
- data/school-selected-subjects.json
- app/src/data/json/school-selected-subjects.json
"""

from __future__ import annotations

import hashlib
import json
import re
from copy import deepcopy
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
EXCEL_FILE = ROOT / "2026학년도 입학생을 위한 2022 개정 교육과정 선택 과목 안내서(EXCEL).xlsm"
SCHOOL_JSON = ROOT / "app" / "src" / "data" / "json" / "school.json"
OUT_SUBJECTS_DATA = ROOT / "data" / "subjects.json"
OUT_SUBJECTS_APP = ROOT / "app" / "src" / "data" / "json" / "subjects.json"
OUT_SCHOOL_DATA = ROOT / "data" / "school-selected-subjects.json"
OUT_SCHOOL_APP = ROOT / "app" / "src" / "data" / "json" / "school-selected-subjects.json"

NORMAL_SHEET = "보통 교과 데이터"
PROFESSIONAL_SHEET = "전문 교과 데이터(목록만 살려)"
CURRICULUM_SHEET = "2022 개정 교육과정"
LIST_SHEET = "목록 데이터"

NORMAL_DATA_LAST_ROW = 110

SECOND_LANGUAGE_ALIASES = {
    "제2외국어": [
        "독일어",
        "프랑스어",
        "스페인어",
        "중국어",
        "일본어",
        "러시아어",
        "아랍어",
        "베트남어",
    ],
    "제2외국어 회화": [
        "독일어 회화",
        "프랑스어 회화",
        "스페인어 회화",
        "중국어 회화",
        "일본어 회화",
        "러시아어 회화",
        "아랍어 회화",
        "베트남어 회화",
    ],
    "심화 제2외국어": [
        "심화 독일어",
        "심화 프랑스어",
        "심화 스페인어",
        "심화 중국어",
        "심화 일본어",
        "심화 러시아어",
        "심화 아랍어",
        "심화 베트남어",
    ],
    "제2외국어권 문화": [
        "독일어권 문화",
        "프랑스어권 문화",
        "스페인어권 문화",
        "중국 문화",
        "일본 문화",
        "러시아 문화",
        "아랍 문화",
        "베트남 문화",
    ],
}

AREA_FALLBACKS = {
    "기술·가정": "기술·가정",
    "로봇과 공학세계": "기술·가정",
    "생활과학 탐구": "기술·가정",
    "창의 공학 설계": "기술·가정",
    "지식 재산 일반": "기술·가정",
    "생애 설계와 자립": "기술·가정",
    "아동발달과 부모": "기술·가정",
    "정보": "정보",
    "인공지능 기초": "정보",
    "데이터 과학": "정보",
    "소프트웨어와 생활": "정보",
    "한문": "한문",
    "한문 고전 읽기": "한문",
    "언어생활과 한자": "한문",
}


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n[ \t]+", "\n", text)
    return text.strip()


def canonical_name(value: Any) -> str:
    text = clean_text(value)
    text = text.replace("ㆍ", "·").replace("⋅", "·").replace("•", "·")
    text = text.replace("・", "·")
    text = text.replace("*", "")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def split_list_text(value: Any) -> list[str]:
    text = clean_text(value)
    if not text:
        return []
    parts: list[str] = []
    for line in text.split("\n"):
        for chunk in line.split(","):
            item = canonical_name(chunk)
            if item:
                parts.append(item)
    return parts


def split_bullets(value: Any) -> list[str]:
    text = clean_text(value)
    if not text:
        return []
    items: list[str] = []
    for line in text.split("\n"):
        item = line.strip()
        item = re.sub(r"^[\u2022\u2023\u2043\u2219\u25e6\u30fb\uff65·•∙ㆍ・\-]\s*", "", item).strip()
        if item:
            items.append(item)
    return items


def build_exploration_activities(data: dict[str, Any]) -> list[dict[str, Any]]:
    activities: list[dict[str, Any]] = []
    for i in range(1, 5):
        task = clean_text(data.get(f"주제탐구과제{i}"))
        examples = unique(split_bullets(data.get(f"활동사례{i}")))
        if task or examples:
            activities.append(
                {
                    "task": task,
                    "activityExamples": examples,
                }
            )
    return activities


def unique(items: list[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for item in items:
        item = canonical_name(item)
        if item and item not in seen:
            seen.add(item)
            result.append(item)
    return result


def normalize_category(value: Any, default: str = "진로선택") -> str:
    text = canonical_name(value).replace(" ", "")
    if text in {"공통", "일반선택", "진로선택", "융합선택"}:
        return text
    if text in {"일반", "일반선택과목"}:
        return "일반선택"
    if text in {"진로", "진로선택과목"} or "전문" in text:
        return "진로선택"
    if text in {"융합", "융합선택과목"}:
        return "융합선택"
    return default


def load_json(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def build_previous_subject_maps() -> tuple[dict[str, str], dict[str, dict[str, Any]]]:
    id_by_name: dict[str, str] = {}
    data_by_name: dict[str, dict[str, Any]] = {}

    for path in (OUT_SUBJECTS_DATA, OUT_SCHOOL_DATA, OUT_SUBJECTS_APP, OUT_SCHOOL_APP):
        data = load_json(path)
        for subject in data.get("subjects", []):
            name = canonical_name(subject.get("name"))
            subject_id = clean_text(subject.get("id"))
            if name and subject_id and name not in id_by_name:
                id_by_name[name] = subject_id
            if name and name not in data_by_name:
                data_by_name[name] = subject

    return id_by_name, data_by_name


def stable_id(name: str, id_by_name: dict[str, str]) -> str:
    if name in id_by_name:
        return id_by_name[name]
    digest = hashlib.sha1(name.encode("utf-8")).hexdigest()[:10]
    return f"subject_{digest}"


def parse_curriculum_area_map(ws: Any) -> dict[str, tuple[str, str]]:
    area_by_name: dict[str, tuple[str, str]] = {}
    current_area = ""
    category_by_col = {3: "일반선택", 4: "진로선택", 5: "융합선택"}

    for row in range(23, 36):
        raw_area = clean_text(ws.cell(row, 1).value)
        if raw_area:
            current_area = canonical_name(raw_area.replace("\n", " "))
            current_area = current_area.replace("사회 (역사/도덕 포함)", "사회")
            current_area = current_area.replace("제2외국어/ 한문", "제2외국어/한문")

        for col, category in category_by_col.items():
            for name in split_list_text(ws.cell(row, col).value):
                area = AREA_FALLBACKS.get(name, current_area)
                if area == "사회 (역사/도덕 포함)":
                    area = "사회"
                if area == "제2외국어/한문":
                    if name in AREA_FALLBACKS:
                        area = AREA_FALLBACKS[name]
                    elif "한문" in name or "한자" in name:
                        area = "한문"
                    else:
                        area = "제2외국어"
                if name in SECOND_LANGUAGE_ALIASES.get("제2외국어", []):
                    area = "제2외국어"
                area_by_name[name] = (area, category)

    for template, aliases in SECOND_LANGUAGE_ALIASES.items():
        if template in area_by_name:
            area, category = area_by_name[template]
            for alias in aliases:
                area_by_name[alias] = ("제2외국어", category)

    return area_by_name


def row_dict(ws: Any, row: int) -> dict[str, Any]:
    headers = [clean_text(ws.cell(1, col).value) for col in range(1, ws.max_column + 1)]
    return {headers[col - 1]: ws.cell(row, col).value for col in range(1, ws.max_column + 1) if headers[col - 1]}


def parse_assessment(data: dict[str, Any]) -> dict[str, Any]:
    text = clean_text(data.get("평가정보"))
    grade_text = clean_text(data.get("석차 등급"))
    achievement = clean_text(data.get("성취도"))
    return {
        "achievementLevel": achievement or None,
        "gradeReported": "미산출" not in text and "5등급" in (text + grade_text),
        "passFail": "성취도(P)" in text or achievement == "P",
        "rawScoreReported": clean_text(data.get("원점수")) == "○",
        "distributionReported": clean_text(data.get("성취도별 분포비율")) == "○",
        "averageReported": clean_text(data.get("과목 평균")) == "○",
        "enrollmentReported": clean_text(data.get("수강자 수")) == "○",
        "sourceText": text,
    }


def default_credits(name: str, previous: dict[str, dict[str, Any]]) -> str:
    prev = previous.get(name, {})
    credits = prev.get("credits")
    if isinstance(credits, str) and credits:
        return credits
    if isinstance(credits, dict):
        default = credits.get("defaultCredits")
        if default:
            return str(default)
    return "3~5"


def normal_subject_from_row(
    ws: Any,
    row: int,
    area_by_name: dict[str, tuple[str, str]],
    id_by_name: dict[str, str],
    previous: dict[str, dict[str, Any]],
    name_override: str | None = None,
    template_name: str | None = None,
) -> dict[str, Any]:
    data = row_dict(ws, row)
    source_name = canonical_name(template_name or data.get("교과목"))
    name = canonical_name(name_override or source_name)
    area, category = area_by_name.get(name) or area_by_name.get(source_name) or ("교양", clean_text(data.get("구분")) or "진로선택")

    content_categories = unique([clean_text(data.get(f"대영역{i}")) for i in range(1, 7)])
    key_ideas = unique([clean_text(data.get(f"생각 열기{i}")) for i in range(1, 7)])
    learning_activities: list[str] = []
    for i in range(1, 7):
        learning_activities.extend(split_bullets(data.get(f"주요 학습활동{i}")))

    interest_fields = unique([clean_text(data.get(f"관심분야{i}")) for i in range(1, 5)])
    exploration_activities = build_exploration_activities(data)
    exploration_tasks = unique([item["task"] for item in exploration_activities])
    activity_examples: list[str] = []
    for item in exploration_activities:
        activity_examples.extend(item["activityExamples"])

    keywords = re.findall(r"#[^\s#]+", clean_text(data.get("주요 키워드")))
    recommended_for = split_bullets(data.get("이런학생 추천"))

    key_contents = unique(content_categories + learning_activities[:8])

    return {
        "id": stable_id(name, id_by_name),
        "name": name,
        "category": normalize_category(category),
        "area": area,
        "credits": default_credits(name, previous),
        "description": clean_text(data.get("과목을 여는 한 줄")),
        "keywords": keywords,
        "keyContents": key_contents,
        "contentCategories": content_categories,
        "keyIdeas": key_ideas,
        "learningActivities": unique(learning_activities),
        "interestFields": interest_fields,
        "relatedDepartments": interest_fields,
        "relatedCareers": [],
        "explorationTasks": exploration_tasks,
        "explorationActivities": exploration_activities,
        "activityExamples": unique(activity_examples),
        "recommendedFor": recommended_for,
        "assessment": parse_assessment(data),
        "suneung": clean_text(data.get("수능 출제 여부")) == "O",
        "source": {
            "type": "excel",
            "file": EXCEL_FILE.name,
            "sheet": NORMAL_SHEET,
            "row": row,
        },
    }


def parse_normal_subjects(
    ws: Any,
    area_by_name: dict[str, tuple[str, str]],
    id_by_name: dict[str, str],
    previous: dict[str, dict[str, Any]],
) -> list[dict[str, Any]]:
    subjects: list[dict[str, Any]] = []
    seen: set[str] = set()

    for row in range(2, NORMAL_DATA_LAST_ROW + 1):
        name = canonical_name(ws.cell(row, 1).value)
        if not name:
            continue
        subject = normal_subject_from_row(ws, row, area_by_name, id_by_name, previous)
        subjects.append(subject)
        seen.add(subject["name"])

        for alias in SECOND_LANGUAGE_ALIASES.get(name, []):
            if alias in seen:
                continue
            alias_subject = normal_subject_from_row(
                ws,
                row,
                area_by_name,
                id_by_name,
                previous,
                name_override=alias,
                template_name=name,
            )
            alias_subject["source"]["templateSubject"] = name
            subjects.append(alias_subject)
            seen.add(alias)

    return subjects


def professional_subject(
    name: str,
    professional_area: str,
    description: str,
    sheet: str,
    row: int,
    id_by_name: dict[str, str],
    previous: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    fallback_description = f"{professional_area or '전문 교과'} 분야의 전문 교과 과목입니다."
    return {
        "id": stable_id(name, id_by_name),
        "name": name,
        "category": "진로선택",
        "area": "전문교과",
        "professionalArea": professional_area,
        "credits": default_credits(name, previous),
        "description": description or fallback_description,
        "keywords": [],
        "keyContents": unique([professional_area] if professional_area else []),
        "relatedCareers": [],
        "relatedDepartments": unique([professional_area] if professional_area else []),
        "curriculumTrack": "전문 교과",
        "suneung": False,
        "source": {
            "type": "excel",
            "file": EXCEL_FILE.name,
            "sheet": sheet,
            "row": row,
        },
    }


def parse_professional_subjects(
    professional_ws: Any,
    list_ws: Any,
    id_by_name: dict[str, str],
    previous: dict[str, dict[str, Any]],
    existing_names: set[str],
) -> list[dict[str, Any]]:
    subjects: list[dict[str, Any]] = []
    seen = set(existing_names)
    current_area = ""

    for row in range(2, professional_ws.max_row + 1):
        name = canonical_name(professional_ws.cell(row, 1).value)
        if not name or name in seen:
            continue
        area = canonical_name(professional_ws.cell(row, 2).value)
        if area:
            current_area = area
        description = clean_text(professional_ws.cell(row, 5).value)
        subject = professional_subject(
            name,
            current_area,
            description,
            PROFESSIONAL_SHEET,
            row,
            id_by_name,
            previous,
        )
        subjects.append(subject)
        seen.add(name)

    for name_col, area_col in ((3, 4), (5, 6)):
        current_area = ""
        for row in range(11, list_ws.max_row + 1):
            name = canonical_name(list_ws.cell(row, name_col).value)
            area = canonical_name(list_ws.cell(row, area_col).value)
            if area:
                current_area = area
            if not name or name.startswith("전문교과") or name in seen:
                continue
            subject = professional_subject(
                name,
                current_area,
                "",
                LIST_SHEET,
                row,
                id_by_name,
                previous,
            )
            subjects.append(subject)
            seen.add(name)

    return subjects


def subject_for_school_only(
    name: str,
    id_by_name: dict[str, str],
    previous: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    prev = deepcopy(previous.get(name, {}))
    if prev:
        prev["id"] = stable_id(name, id_by_name)
        prev["name"] = name
        prev["category"] = normalize_category(prev.get("category"), default="융합선택")
        prev["credits"] = default_credits(name, previous)
        if prev.get("area") in {None, "", "고시 외 과목", "기타"}:
            prev["area"] = "교양"
        prev.setdefault("source", {})
        prev["sourceStatus"] = "supplemented_from_previous_data"
        prev["sourceNote"] = "새 Excel 원본에서 찾지 못해 기존 보강 데이터를 보존했습니다."
        return prev
    return {
        "id": stable_id(name, id_by_name),
        "name": name,
        "category": "진로선택",
        "area": "기타",
        "credits": "3",
        "description": f"{name} 과목 정보는 학교 편성표에는 있으나 새 Excel 원본에서 찾지 못했습니다.",
        "keyContents": [],
        "relatedCareers": [],
        "relatedDepartments": [],
        "suneung": False,
        "sourceStatus": "school_only",
    }


def expand_subject_name(name: str) -> list[str]:
    if "↔" not in name:
        return [canonical_name(name)]
    return [canonical_name(part) for part in name.split("↔") if canonical_name(part)]


def collect_school_selection_offerings(school: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    offerings_by_name: dict[str, list[dict[str, Any]]] = {}
    for cohort, cohort_data in school.get("cohorts", {}).items():
        for group in cohort_data.get("selections", []):
            for option in group.get("options", []):
                for name in expand_subject_name(option):
                    offerings_by_name.setdefault(name, []).append(
                        {
                            "cohort": cohort,
                            "grade": group.get("grade"),
                            "semester": group.get("semester"),
                            "groupId": group.get("id"),
                            "groupLabel": group.get("label"),
                            "choose": group.get("choose"),
                            "creditsEach": group.get("creditsEach"),
                            "totalCredits": group.get("totalCredits"),
                        }
                    )
    return offerings_by_name


def school_selected_subject(
    subject: dict[str, Any],
    offerings: list[dict[str, Any]],
) -> dict[str, Any]:
    selected = deepcopy(subject)
    selected["sourceStatus"] = selected.get("sourceStatus", "matched")
    credits = selected.get("credits")
    if isinstance(credits, str):
        selected["credits"] = {
            "sourceCreditText": credits,
            "defaultCredits": None,
        }
    selected["offerings"] = offerings
    return selected


def write_json(path: Path, data: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main() -> None:
    if not EXCEL_FILE.exists():
        raise FileNotFoundError(EXCEL_FILE)

    id_by_name, previous = build_previous_subject_maps()
    workbook = load_workbook(EXCEL_FILE, data_only=True, keep_vba=True)

    area_by_name = parse_curriculum_area_map(workbook[CURRICULUM_SHEET])
    normal_subjects = parse_normal_subjects(
        workbook[NORMAL_SHEET],
        area_by_name,
        id_by_name,
        previous,
    )

    all_subjects = normal_subjects[:]
    existing_names = {subject["name"] for subject in all_subjects}
    professional_subjects = parse_professional_subjects(
        workbook[PROFESSIONAL_SHEET],
        workbook[LIST_SHEET],
        id_by_name,
        previous,
        existing_names,
    )
    all_subjects.extend(professional_subjects)

    school = load_json(SCHOOL_JSON)
    offerings_by_name = collect_school_selection_offerings(school)

    by_name = {subject["name"]: subject for subject in all_subjects}
    missing_school_names = sorted(name for name in offerings_by_name if name not in by_name)
    for name in missing_school_names:
        supplemental = subject_for_school_only(name, id_by_name, previous)
        by_name[name] = supplemental
        all_subjects.append(supplemental)

    subjects_payload = {
        "metadata": {
            "generatedAt": date.today().isoformat(),
            "source": EXCEL_FILE.name,
            "normalSheet": NORMAL_SHEET,
            "professionalSheet": PROFESSIONAL_SHEET,
            "professionalListFallbackSheet": LIST_SHEET,
            "normalSubjectCount": len(normal_subjects),
            "professionalSubjectCount": len(professional_subjects),
            "schoolOnlySupplementCount": len(missing_school_names),
            "schoolOnlySupplementSubjects": missing_school_names,
            "subjectCount": len(all_subjects),
        },
        "subjects": all_subjects,
    }

    school_subjects = [
        school_selected_subject(by_name[name], offerings)
        for name, offerings in sorted(offerings_by_name.items())
    ]
    matched = [s for s in school_subjects if s.get("sourceStatus") == "matched"]
    supplemented = [s for s in school_subjects if s.get("sourceStatus") != "matched"]
    school_payload = {
        "metadata": {
            "generatedAt": date.today().isoformat(),
            "schoolSource": str(SCHOOL_JSON.relative_to(ROOT)).replace("\\", "/"),
            "excelSource": EXCEL_FILE.name,
            "normalSheet": NORMAL_SHEET,
            "professionalSheet": PROFESSIONAL_SHEET,
            "professionalListFallbackSheet": LIST_SHEET,
            "scope": "school_selection_options",
            "subjectCount": len(school_subjects),
            "matchedCount": len(matched),
            "supplementedCount": len(supplemented),
            "supplementedSubjects": [s["name"] for s in supplemented],
            "unresolvedCount": 0,
            "unresolvedSubjects": [],
        },
        "subjects": school_subjects,
    }

    write_json(OUT_SUBJECTS_DATA, subjects_payload)
    write_json(OUT_SUBJECTS_APP, subjects_payload)
    write_json(OUT_SCHOOL_DATA, school_payload)
    write_json(OUT_SCHOOL_APP, school_payload)

    print(f"subjects: {len(all_subjects)}")
    print(f"school-selected subjects: {len(school_subjects)}")
    print(f"school-only supplements: {', '.join(missing_school_names) if missing_school_names else 'none'}")


if __name__ == "__main__":
    main()
