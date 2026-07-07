# -*- coding: utf-8 -*-
"""
대교협 「2028학년도 권역별 대학별 권장과목」(2026-02-20) 파서.

입력:  260220-2028학년도 권역별 대학별 권장과목.xlsx (Sheet1, 5행부터 데이터)
출력:  data/university-recommendations.json
       app/src/data/json/university-recommendations.json

셀 서식이 대학마다 제각각(중첩 괄호, 줄바꿈, 오탈자, 서술형)이라
구분자 분리 대신 과목 사전 최장일치 스캔으로 과목명을 추출한다.
사전은 data/subjects.json의 과목명 + 교과(군) 우산 용어로 구성한다.
"""
import json
import os
import re
from collections import Counter

import openpyxl

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "data")
APP_JSON_DIR = os.path.join(BASE_DIR, "app", "src", "data", "json")
SOURCE_XLSX = os.path.join(BASE_DIR, "260220-2028학년도 권역별 대학별 권장과목.xlsx")

# 교과(군) 우산 용어 — 특정 과목이 아니라 영역 전체를 가리키는 표기
AREA_TERMS = [
    "국어", "수학", "영어", "사회", "과학", "한국사", "한문",
    "제2외국어", "정보", "교양", "기술·가정", "예술", "체육",
    "역사", "도덕", "윤리", "지리", "미적분",
]

# 원본 표기 변형/오탈자 → 정식 명칭 별칭
ALIASES = {
    "물리": "물리학",
    "생물과 유전": "생물의 유전",
    "물리과 에너지": "역학과 에너지",
    "사화와 문화": "사회와 문화",
}

# 미매칭 리포트에서 제외할 구조 라벨/조사류
REPORT_STOPWORDS = {
    "일반선택", "진로선택", "융합선택", "교과", "과목", "포함", "또는",
    "관련과목", "이수", "권장", "일반", "교과영역이수권장", "선택",
}

# 서술형(계열 자율) 판정 마커
FLEXIBLE_MARKERS = [
    "고려하여", "자유롭게", "구분 없이", "제한 없", "상관없이",
    "희망 진로에 따라", "적성에 따라", "적성에 맞", "특성에 맞",
]


def norm_for_match(text: str) -> str:
    """매칭용 정규화: 공백/가운뎃점 제거, 로마숫자 통일."""
    t = text.replace("․", "·").replace("･", "·")
    t = re.sub(r"[\s·]+", "", t)
    t = t.replace("II", "Ⅱ").replace("I", "Ⅰ")
    t = t.replace("2", "Ⅱ").replace("1", "Ⅰ") if re.search(r"[가-힣][12]$", t) else t
    return t


def load_vocabulary() -> dict:
    """정규화명 → 정식 과목명 사전."""
    with open(os.path.join(DATA_DIR, "subjects.json"), encoding="utf-8") as f:
        data = json.load(f)
    subjects = data if isinstance(data, list) else data.get("subjects", [])
    vocab = {}
    for s in subjects:
        name = s.get("name", "").strip()
        if len(name) >= 2:
            vocab[norm_for_match(name)] = name
    # 우산 용어는 과목명과 충돌하지 않을 때만 추가
    for term in AREA_TERMS:
        vocab.setdefault(norm_for_match(term), term)
    for variant, canonical in ALIASES.items():
        vocab.setdefault(norm_for_match(variant), canonical)
    return vocab


def scan_subjects(cell: str, vocab: dict, keys_by_len: list) -> tuple:
    """최장일치 스캔. (특정 과목 목록, 우산 용어 목록, 미매칭 한글 토큰) 반환."""
    normalized = norm_for_match(cell)
    n = len(normalized)
    consumed = [False] * n
    found = []
    i = 0
    while i < n:
        matched = False
        for key in keys_by_len:
            if normalized.startswith(key, i):
                found.append(vocab[key])
                for j in range(i, i + len(key)):
                    consumed[j] = True
                i += len(key)
                matched = True
                break
        if not matched:
            i += 1
    leftover = "".join(
        c if not consumed[k] else " " for k, c in enumerate(normalized)
    )
    unmatched = [
        t for t in re.findall(r"[가-힣Ⅰ-Ⅻ]{2,}", leftover)
        if t not in REPORT_STOPWORDS
    ]
    area_set = set(AREA_TERMS)
    specific, areas = [], []
    for name in found:
        target = areas if name in area_set else specific
        if name not in target:
            target.append(name)
    return specific, areas, unmatched


def parse_cell(cell, vocab, keys_by_len, unmatched_counter):
    raw = str(cell).strip() if cell is not None else ""
    if not raw or raw == "-":
        return None
    subjects, areas, unmatched = scan_subjects(raw, vocab, keys_by_len)
    for t in unmatched:
        unmatched_counter[t] += 1
    return {
        "raw": raw,
        "subjects": subjects,
        "areas": areas,
        "isFlexible": any(m in raw for m in FLEXIBLE_MARKERS),
    }


def clean(text) -> str:
    if text is None:
        return ""
    return re.sub(r"\s+", " ", str(text)).strip()


def main():
    vocab = load_vocabulary()
    keys_by_len = sorted(vocab.keys(), key=len, reverse=True)
    unmatched_counter = Counter()

    wb = openpyxl.load_workbook(SOURCE_XLSX, read_only=True)
    ws = wb["Sheet1"]
    legend = clean(ws.cell(2, 1).value) if not isinstance(ws, type(None)) else ""

    entries = []
    for row in ws.iter_rows(min_row=5, max_col=8, values_only=True):
        region, district, uni, unit_group, unit_detail, core, rec, note = (
            list(row) + [None] * 8
        )[:8]
        uni_name = clean(uni)
        if not uni_name:
            continue
        entry = {
            "region": clean(region),
            "district": clean(district),
            "university": uni_name,
            "unitGroup": clean(unit_group),
            "unit": clean(unit_detail) or clean(unit_group),
            "core": parse_cell(core, vocab, keys_by_len, unmatched_counter),
            "recommended": parse_cell(rec, vocab, keys_by_len, unmatched_counter),
            "note": clean(note) if clean(note) != "-" else "",
        }
        entries.append(entry)
    wb.close()

    result = {
        "title": "2028학년도 권역별 대학별 권장과목",
        "legend": legend,
        "source": {
            "publisher": "한국대학교육협의회(대교협) 대입상담센터",
            "document": "2028학년도 권역별 대학별 권장과목",
            "published": "2026-02-20",
            "file": os.path.basename(SOURCE_XLSX),
            "basis": "각 대학이 발표한 2028학년도 모집단위별 반영과목(2025-09-30 어디가 탑재)을 대교협이 취합한 참고자료",
            "note": "핵심과목=필수적 이수를 권장하는 과목, 권장과목=가급적 이수를 권장하는 과목. 필수 이수 기준이 아니며 수시 갱신됨.",
        },
        "entries": entries,
    }

    for out_dir in (DATA_DIR, APP_JSON_DIR):
        path = os.path.join(out_dir, "university-recommendations.json")
        with open(path, "w", encoding="utf-8") as f:
            json.dump(result, f, ensure_ascii=False, indent=2)
        print(f"wrote {path}")

    unis = {e["university"] for e in entries}
    with_core = sum(1 for e in entries if e["core"] and e["core"]["subjects"])
    flexible = sum(1 for e in entries if e["core"] and e["core"]["isFlexible"])
    print(f"entries={len(entries)} universities={len(unis)} "
          f"core_with_subjects={with_core} flexible={flexible}")
    print("unmatched top20:", unmatched_counter.most_common(20))


if __name__ == "__main__":
    main()
