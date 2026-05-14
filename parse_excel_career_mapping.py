from __future__ import annotations

import argparse
import json
import re
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


SOURCE_FILE_KEYWORD = "2022 개정 교육과정 선택 과목 안내서"
TRACK_SHEET_INDEX = 12
DEPARTMENT_SHEET_INDEX = 13

CATEGORIES = ("일반선택", "진로선택", "융합선택")

FIELD_META: dict[str, str] = {
    "humanities": "인문 분야",
    "social_sciences": "사회 분야",
    "natural_sciences": "자연 분야",
    "engineering": "공학 분야",
    "health_medicine": "보건·의약학 분야",
    "education": "교육 분야",
    "arts_sports": "예술·체육 분야",
    "interdisciplinary": "자율전공 분야",
}

TRACK_META: dict[str, tuple[str, str, str, str]] = {
    "언어·문학": ("humanities", "language_literature", "언어·문학 계열", "인문 분야"),
    "인문과학": ("humanities", "humanities_science", "인문과학 계열", "인문 분야"),
    "상경": ("social_sciences", "business_economics", "상경 계열", "사회 분야"),
    "법·행정": ("social_sciences", "law_administration", "법·행정 계열", "사회 분야"),
    "광고언론정보": ("social_sciences", "media_communication", "광고·언론·정보 계열", "사회 분야"),
    "사회과학": ("social_sciences", "social_science", "사회과학 계열", "사회 분야"),
    "자연과학": ("natural_sciences", "natural_science", "자연과학 계열", "자연 분야"),
    "생활과학": ("natural_sciences", "life_science_track", "생활과학 계열", "자연 분야"),
    "농학": ("natural_sciences", "agriculture", "농학 계열", "자연 분야"),
    "기계·전기·전자": ("engineering", "mechanical_electrical", "기계·전기·전자 계열", "공학 분야"),
    "건축·환경": ("engineering", "architecture_environment", "건축·환경 계열", "공학 분야"),
    "화학·생명": ("engineering", "chemical_bio", "화학·생명 계열", "공학 분야"),
    "정보·컴퓨터": ("engineering", "it_software", "정보·컴퓨터 계열", "공학 분야"),
    "보건": ("health_medicine", "health", "보건 계열", "보건·의약학 분야"),
    "의약학": ("health_medicine", "medicine_pharmacy", "의약학 계열", "보건·의약학 분야"),
    "교육": ("education", "education", "교육 계열", "교육 분야"),
    "체육": ("arts_sports", "physical_education", "체육 계열", "예술·체육 분야"),
    "예술": ("arts_sports", "arts", "예술 계열", "예술·체육 분야"),
    "자율전공": ("interdisciplinary", "interdisciplinary", "자율전공 분야", "자율전공 분야"),
}

TRACK_ORDER = list(TRACK_META)
FIELD_ORDER = list(FIELD_META)

SUBJECT_NAME_FIXES = {
    "영어 독해과 작문": "영어 독해와 작문",
    "영어I": "영어Ⅰ",
    "영어II": "영어Ⅱ",
    "미적분I": "미적분Ⅰ",
    "미적분II": "미적분Ⅱ",
    "동아시아 역사기행": "동아시아 역사 기행",
    "로봇과 공학 세계": "로봇과 공학세계",
    "제2외국어 문화": "제2외국어권 문화",
    "제2외국어 회화 과목": "제2외국어 회화",
    "독서토론과 글쓰기": "독서 토론과 글쓰기",
    "심화중국어": "심화 중국어",
    "기후변화와 생태": "기후변화와 환경생태",
    "매체와 의사소통": "매체 의사소통",
}

SUBJECT_SPLIT_FIXES = {
    "전자기와 양자 데이터 과학": ["전자기와 양자", "데이터 과학"],
}

NON_SUBJECT_LABELS = {
    "인문",
    "사회 계열 교육과",
}


def clean_cell(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text == "#VALUE!":
        return ""
    return re.sub(r"[ \t]+", " ", text.replace("\r\n", "\n").replace("\r", "\n")).strip()


def strip_trailing_etc(text: str) -> str:
    text = re.sub(r"\s*[,.·ㆍ]*\s*등[.)]*\s*$", "", text.strip())
    return text.strip(" ,.;")


def split_comma_list(value: Any) -> list[str]:
    text = clean_cell(value)
    if not text:
        return []
    text = text.replace("，", ",")
    items: list[str] = []
    for item in re.split(r"[,;\n]+", text):
        item = re.sub(r"^[•∙ㆍ\-\s]+", "", item).strip()
        item = strip_trailing_etc(item)
        if item and item not in {"등", "해당 없음"}:
            items.append(item)
    return dedupe(items)


def split_bullets(value: Any) -> list[str]:
    text = clean_cell(value)
    if not text:
        return []
    parts = re.split(r"(?:^|\n|\s{2,})[•∙ㆍ]\s*", text)
    if len(parts) <= 1:
        parts = re.split(r"\n+", text)
    cleaned = []
    for part in parts:
        item = re.sub(r"^[•∙ㆍ\-\s]+", "", part).strip()
        item = strip_trailing_etc(item)
        if item:
            cleaned.append(item)
    return dedupe(cleaned)


def dedupe(items: list[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for item in items:
        if item in seen:
            continue
        seen.add(item)
        result.append(item)
    return result


def find_excel_file() -> Path:
    candidates = [
        path
        for path in Path(".").glob("*.xlsm")
        if SOURCE_FILE_KEYWORD in path.name or "선택 과목 안내서" in path.name
    ]
    if not candidates:
        raise FileNotFoundError("선택 과목 안내서 xlsm 파일을 찾지 못했습니다.")
    return candidates[0]


def load_subject_names() -> set[str]:
    data = json.loads(Path("app/src/data/json/subjects.json").read_text(encoding="utf-8"))
    subjects = data["subjects"] if isinstance(data, dict) else data
    return {subject["name"] for subject in subjects}


def canonical_subject_name(name: str, subject_names: set[str]) -> str | None:
    name = strip_trailing_etc(name)
    name = SUBJECT_NAME_FIXES.get(name, name)
    if name in subject_names:
        return name

    compact = re.sub(r"\s+", "", name)
    compact_index = {re.sub(r"\s+", "", subject): subject for subject in subject_names}
    if compact in compact_index:
        return compact_index[compact]

    return None


def split_subjects(value: Any, subject_names: set[str], unknown: dict[str, int]) -> list[str]:
    text = clean_cell(value)
    if not text:
        return []

    # Subject columns occasionally prefix a list with labels such as
    # "인문, 사회 계열 교육과:". The label is not a subject.
    text = re.sub(r"(^|[\n;])\s*[^,\n:]{1,40}:\s*", r"\1", text)
    text = text.replace("，", ",")

    subjects: list[str] = []
    for raw in re.split(r"[,;\n]+", text):
        item = re.sub(r"^[•∙ㆍ\-\s]+", "", raw).strip()
        if ":" in item:
            item = item.split(":", 1)[1].strip()
        item = strip_trailing_etc(item)
        item = SUBJECT_NAME_FIXES.get(item, item)
        if item in SUBJECT_SPLIT_FIXES:
            for split_item in SUBJECT_SPLIT_FIXES[item]:
                subject = canonical_subject_name(split_item, subject_names)
                if subject:
                    subjects.append(subject)
            continue
        if (
            not item
            or item in {"등", "해당 없음"}
            or item in NON_SUBJECT_LABELS
            or item.startswith("※")
            or "전 과목" in item
            or "필요" in item
            or "이수해야" in item
        ):
            continue
        subject = canonical_subject_name(item, subject_names)
        if subject:
            subjects.append(subject)
        else:
            unknown[item] += 1
    return dedupe(subjects)


def parse_universities(row: tuple[Any, ...]) -> dict[str, list[str]]:
    return {
        "서울": split_comma_list(row[10] or row[7]),
        "수도권": split_comma_list(row[11] or row[8]),
        "비수도권": split_comma_list(row[12] or row[9]),
    }


def build_major_courses(base: Any, advanced: Any) -> dict[str, str]:
    result: dict[str, str] = {}
    base_text = clean_cell(base)
    advanced_text = clean_cell(advanced)
    if base_text:
        result["전공 기초"] = base_text
    if advanced_text:
        result["전공 심화"] = advanced_text
    return result


def parse_workbook(path: Path) -> tuple[dict[str, Any], dict[str, int]]:
    subject_names = load_subject_names()
    unknown_subjects: dict[str, int] = defaultdict(int)

    workbook = load_workbook(path, data_only=True, read_only=True, keep_vba=True)
    track_sheet = workbook.worksheets[TRACK_SHEET_INDEX]
    dept_sheet = workbook.worksheets[DEPARTMENT_SHEET_INDEX]

    track_dept_map: dict[str, str] = {}
    tracks: dict[str, dict[str, Any]] = {}

    for row_index in range(2, track_sheet.max_row + 1):
        track_name = clean_cell(track_sheet.cell(row_index, 1).value)
        if not track_name or track_name not in TRACK_META:
            continue

        field_id, track_id, display_name, _ = TRACK_META[track_name]
        tracks[track_name] = {
            "id": track_id,
            "name": display_name,
            "sourceName": track_name,
            "recommendedSubjects": {
                "일반선택": split_subjects(track_sheet.cell(row_index, 3).value, subject_names, unknown_subjects),
                "진로선택": split_subjects(track_sheet.cell(row_index, 4).value, subject_names, unknown_subjects),
                "융합선택": split_subjects(track_sheet.cell(row_index, 5).value, subject_names, unknown_subjects),
            },
            "tip": clean_cell(track_sheet.cell(row_index, 6).value),
            "departments": [],
            "fieldId": field_id,
        }

        for dept_name in split_comma_list(track_sheet.cell(row_index, 2).value):
            track_dept_map[dept_name] = track_name

    for row_index in range(2, dept_sheet.max_row + 1):
        row = tuple(dept_sheet.cell(row_index, col).value for col in range(1, 18))
        dept_name = clean_cell(row[0])
        if not dept_name:
            continue

        has_detail = any(clean_cell(value) for value in row[1:])
        if not has_detail:
            continue

        track_name = track_dept_map.get(dept_name)
        if not track_name:
            if dept_name in {"문예창작학과", "무용학과"}:
                track_name = "예술"
            else:
                print(f"[warn] 계열 데이터에 없는 학과: {dept_name} (row {row_index})")
                continue

        department = {
            "name": dept_name,
            "description": clean_cell(row[1]),
            "recommendedStudents": split_bullets(row[4]),
            "recommendedSubjects": {
                "일반선택": split_subjects(row[14], subject_names, unknown_subjects),
                "진로선택": split_subjects(row[15], subject_names, unknown_subjects),
                "융합선택": split_subjects(row[16], subject_names, unknown_subjects),
            },
            "majorCourses": build_major_courses(row[2], row[3]),
            "similarDepartments": split_comma_list(row[5]),
            "universities": parse_universities(row),
            "careers": split_bullets(row[13]),
            "source": {
                "type": "excel",
                "file": path.name,
                "sheet": dept_sheet.title,
                "row": row_index,
            },
        }
        tracks[track_name]["departments"].append(department)

    fields = [
        {
            "id": field_id,
            "name": field_name,
            "tracks": [
                {
                    key: value
                    for key, value in tracks[track_name].items()
                    if key != "fieldId"
                }
                for track_name in TRACK_ORDER
                if track_name in tracks and tracks[track_name]["fieldId"] == field_id
            ],
        }
        for field_id, field_name in FIELD_META.items()
    ]

    result = {
        "metadata": {
            "sourceFile": path.name,
            "sourceSheets": [track_sheet.title, dept_sheet.title],
            "generatedAt": datetime.now().isoformat(timespec="seconds"),
            "fieldCount": len(fields),
            "trackCount": sum(len(field["tracks"]) for field in fields),
            "departmentCount": sum(
                len(track["departments"]) for field in fields for track in field["tracks"]
            ),
        },
        "fields": fields,
    }
    return result, dict(sorted(unknown_subjects.items(), key=lambda item: (-item[1], item[0])))


def backup_existing_files() -> Path:
    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup_dir = Path("data/backups") / f"career-mapping-excel-reparse-{timestamp}"
    backup_dir.mkdir(parents=True, exist_ok=False)
    for source, name in [
        (Path("data/career-mapping.json"), "data-career-mapping.json"),
        (Path("app/src/data/json/career-mapping.json"), "app-career-mapping.json"),
    ]:
        if source.exists():
            shutil.copy2(source, backup_dir / name)
    return backup_dir


def write_json(path: Path, data: dict[str, Any]) -> None:
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="Parse Excel career mapping data.")
    parser.add_argument("--no-backup", action="store_true", help="Skip backup before writing JSON files.")
    args = parser.parse_args()

    excel_path = find_excel_file()
    data, unknown_subjects = parse_workbook(excel_path)
    backup_dir = None if args.no_backup else backup_existing_files()

    write_json(Path("data/career-mapping.json"), data)
    write_json(Path("app/src/data/json/career-mapping.json"), data)

    if backup_dir:
        print(f"Backup: {backup_dir}")
    print(f"Source: {excel_path}")
    print(
        "Generated: "
        f"{data['metadata']['fieldCount']} fields, "
        f"{data['metadata']['trackCount']} tracks, "
        f"{data['metadata']['departmentCount']} departments"
    )
    if unknown_subjects:
        print("Unknown subjects ignored:")
        for name, count in list(unknown_subjects.items())[:50]:
            print(f"  {name}: {count}")


if __name__ == "__main__":
    main()
